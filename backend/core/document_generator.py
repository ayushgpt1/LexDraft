from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
import re


def build_affidavit_docx(case_data, generated_paragraphs, output_path):

    doc = Document()

    # =========================================================
    # PAGE SETUP
    # =========================================================

    section = doc.sections[0]

    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # =========================================================
    # DEFAULT FONT
    # =========================================================

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # =========================================================
    # HELPER FUNCTIONS
    # =========================================================

    def format_paragraph(
        p,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=0,
        after=0,
        line_spacing=1.0,
        left_indent=None,
        first_line_indent=None
    ):

        p.alignment = alignment

        pf = p.paragraph_format

        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line_spacing

        if left_indent is not None:
            pf.left_indent = Inches(left_indent)

        if first_line_indent is not None:
            pf.first_line_indent = Inches(first_line_indent)

        return p


    def format_run(
        run,
        bold=False,
        italic=False
    ):

        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = bold
        run.italic = italic

        return run


    def add_paragraph(
        text="",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        bold=False,
        before=0,
        after=0,
        line_spacing=1.0,
        left_indent=None,
        first_line_indent=None
    ):

        p = doc.add_paragraph()

        format_paragraph(
            p,
            alignment=alignment,
            before=before,
            after=after,
            line_spacing=line_spacing,
            left_indent=left_indent,
            first_line_indent=first_line_indent
        )

        run = p.add_run(text)

        format_run(
            run,
            bold=bold
        )

        return p


    def add_centered(
        text,
        bold=False,
        before=0,
        after=0
    ):

        return add_paragraph(
            text=text,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=bold,
            before=before,
            after=after,
            line_spacing=1.0
        )


    def add_justified(
        text,
        before=0,
        after=0,
        line_spacing=1.25,
        left_indent=None,
        first_line_indent=None
    ):

        return add_paragraph(
            text=text,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            before=before,
            after=after,
            line_spacing=line_spacing,
            left_indent=left_indent,
            first_line_indent=first_line_indent
        )


    # =========================================================
    # DATE FORMATTER
    # =========================================================

    def format_legal_date(date_string):

        parts = date_string.split()

        if len(parts) >= 3:

            try:

                day = int(parts[0])

                if 10 <= day % 100 <= 20:
                    suffix = "th"

                else:
                    suffix = {
                        1: "st",
                        2: "nd",
                        3: "rd"
                    }.get(day % 10, "th")

                parts[0] = f"{day}{suffix}"

            except ValueError:
                pass

        formatted = " ".join(parts)

        # Sample-style wording:
        # 5 September 2026
        # ->
        # 5th day of September 2026

        if " day of " not in formatted:
            formatted_parts = formatted.split()

            if len(formatted_parts) >= 3:
                formatted = (
                    f"{formatted_parts[0]} day of "
                    f"{formatted_parts[1]} "
                    f"{formatted_parts[2]}"
                )

        return formatted


    formatted_date = format_legal_date(
        case_data.date
    )


    # =========================================================
    # 1. COURT HEADING
    # =========================================================

    add_centered(
        case_data.court,
        bold=True,
        before=0,
        after=0
    )

    add_centered(
        case_data.jurisdiction,
        bold=True,
        before=0,
        after=0
    )

    # FIX #1:
    # Increased spacing AFTER case number
    add_centered(
        f"{case_data.proceeding_type} NO. "
        f"{case_data.case_number} OF {case_data.year}",
        bold=True,
        before=0,
        after=18
    )


    # =========================================================
    # 2. CAUSE TITLE
    # =========================================================

    # -------------------------
    # Petitioner
    # -------------------------

    p = doc.add_paragraph()

    format_paragraph(
        p,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=0,
        after=0,
        line_spacing=1.0
    )

    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5),
        WD_TAB_ALIGNMENT.RIGHT
    )

    run = p.add_run(
        case_data.petitioner
    )

    format_run(run)

    run = p.add_run(
        "\t...Petitioner"
    )

    format_run(run)


    # -------------------------
    # VERSUS
    # -------------------------

    # FIX #1:
    # Increased spacing AFTER VERSUS
    add_centered(
        "VERSUS",
        bold=True,
        before=12,
        after=12
    )


    # -------------------------
    # Respondents
    # -------------------------

    for respondent in case_data.respondents:

        p = doc.add_paragraph()

        format_paragraph(
            p,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            before=0,
            after=0,
            line_spacing=1.0
        )

        p.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.5),
            WD_TAB_ALIGNMENT.RIGHT
        )

        run = p.add_run(
            f"{respondent.number}. "
            f"{respondent.name}"
        )

        format_run(run)

        run = p.add_run(
            f"\t...Respondent No.{respondent.number}"
        )

        format_run(run)


    # =========================================================
    # 3. AFFIDAVIT TITLE
    # =========================================================

    add_centered(
        f"AFFIDAVIT IN REPLY ON BEHALF OF "
        f"RESPONDENT NO. "
        f"{case_data.answering_respondent_number}",
        bold=True,
        before=12,
        after=12
    )


    # =========================================================
    # 4. DEPONENT CLAUSE
    # =========================================================

    deponent = case_data.deponent

    deponent_text = (
        f"I, {deponent.name}, "
        f"{deponent.designation}, "
        f"residing at {deponent.address}, "
        f"the Respondent No."
        f"{case_data.answering_respondent_number} "
        f"above named, do hereby solemnly affirm "
        f"and state as under:"
    )

    # FIX #2:
    # Same left indentation as numbered paragraphs
    add_justified(
        deponent_text,
        before=0,
        after=12,
        line_spacing=1.25,
        left_indent=0.115
    )


    # =========================================================
    # 5. NUMBERED PARAGRAPHS
    # =========================================================

    for paragraph in generated_paragraphs.paragraphs:

        p = doc.add_paragraph()

        format_paragraph(
            p,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            before=8,
            after=0,
            line_spacing=1.25,
            left_indent=0.115
        )

        # Paragraph number
        number_run = p.add_run(
            f"{paragraph.paragraph_number}. "
        )

        format_run(
            number_run,
            bold=True
        )

        content = paragraph.content


        # =====================================================
        # EXHIBIT DETECTION
        # =====================================================

        exhibit_pattern = (
            r"EXHIBIT"
            r"[-‐-‒–—]"
            r"\s*"
            r"[‘']A[’']"
        )

        match = re.search(
            exhibit_pattern,
            content,
            flags=re.IGNORECASE
        )


        if match:

            before_exhibit = content[
                :match.start()
            ]

            exhibit_text = content[
                match.start():match.end()
            ]

            after_exhibit = content[
                match.end():]


            if before_exhibit:

                run = p.add_run(
                    before_exhibit
                )

                format_run(run)


            # EXHIBIT-‘A’ BOLD
            run = p.add_run(
                exhibit_text
            )

            format_run(
                run,
                bold=True
            )


            if after_exhibit:

                run = p.add_run(
                    after_exhibit
                )

                format_run(run)

        else:

            run = p.add_run(
                content
            )

            format_run(run)


    # =========================================================
    # 6. PAGE BREAK BEFORE PRAYER
    # =========================================================

    doc.add_page_break()


    # =========================================================
    # 7. PRAYER
    # =========================================================

    add_centered(
        "PRAYER",
        bold=True,
        before=0,
        after=13
    )


    add_justified(
        "I therefore respectfully pray that this Hon’ble Court "
        "may be pleased to:",
        before=13,
        after=0,
        line_spacing=1.0
    )


    prayer_items = [

        (
            "dismiss the present "
            f"{case_data.proceeding_type.title()} "
            "with costs;"
        ),

        (
            "refuse any interim or ad-interim relief "
            "sought by the Petitioner; and"
        ),

        (
            "grant such other and further reliefs as this "
            "Hon’ble Court may deem fit and proper in the "
            "facts and circumstances of the case."
        )

    ]


    for index, prayer_text in enumerate(
        prayer_items,
        start=1
    ):

        letter = chr(
            96 + index
        )

        p = doc.add_paragraph()

        format_paragraph(
            p,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            before=11.5,
            after=0,
            line_spacing=1.0,
            left_indent=0.63,
            first_line_indent=-0.235
        )

        run = p.add_run(
            f"({letter}) "
        )

        format_run(
            run,
            bold=True
        )

        run = p.add_run(
            prayer_text
        )

        format_run(run)


    # =========================================================
    # 8. JURAT
    # =========================================================

    add_paragraph(
        f"Solemnly affirmed at "
        f"{case_data.place}",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=18,
        after=0,
        line_spacing=1.0
    )


    # FIX #3:
    # "On this 5th day of September 2026"
    add_paragraph(
        f"On this {formatted_date}",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=0,
        after=10,
        line_spacing=1.0
    )


    # =========================================================
    # BEFORE ME + DEPONENT ON SAME LINE
    # =========================================================

    # FIX #3:
    # Before Me and DEPONENT on SAME horizontal line

    p = doc.add_paragraph()

    format_paragraph(
        p,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=0,
        after=12,
        line_spacing=1.0
    )

    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5),
        WD_TAB_ALIGNMENT.RIGHT
    )

    run = p.add_run(
        "Before Me"
    )

    format_run(run)

    run = p.add_run(
        "\tDEPONENT"
    )

    format_run(
        run,
        bold=True
    )


    # =========================================================
    # 9. VERIFICATION
    # =========================================================

    add_centered(
        "VERIFICATION",
        bold=True,
        before=12,
        after=13
    )


    paragraph_count = len(
        generated_paragraphs.paragraphs
    )


    verification_text = (
        f"I, {deponent.name}, the Deponent above named, "
        f"do hereby verify that the contents of paragraphs "
        f"1 to {paragraph_count} and the Prayer above are true "
        f"and correct to my knowledge and belief and that "
        f"nothing material has been concealed therefrom."
    )


    add_justified(
        verification_text,
        before=13,
        after=0,
        line_spacing=1.25
    )


    # FIX #3:
    # Uses "5th day of September 2026"
    add_justified(
        f"Verified at {case_data.place} "
        f"on this {formatted_date}.",
        before=10,
        after=0,
        line_spacing=1.0
    )


    # DEPONENT
    p = doc.add_paragraph()

    format_paragraph(
        p,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        before=12,
        after=12,
        line_spacing=1.0
    )

    run = p.add_run(
        "DEPONENT"
    )

    format_run(
        run,
        bold=True
    )


    # =========================================================
    # 10. ADVOCATE BLOCK
    # =========================================================

    p = doc.add_paragraph()

    format_paragraph(
        p,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        before=0,
        after=0,
        line_spacing=1.0,
        left_indent=0.115
    )

    run = p.add_run(
        case_data.advocate_firm
    )

    format_run(
        run,
        bold=True
    )


    add_justified(
        f"Advocate for {case_data.advocate_for}",
        before=5.5,
        after=0,
        line_spacing=1.0
    )


    # =========================================================
    # SAVE
    # =========================================================

    doc.save(
        output_path
    )

    print(
        f"Document created successfully: {output_path}"
    )
def extract_docx_text(file_path):
    document = Document(file_path)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(paragraphs) 